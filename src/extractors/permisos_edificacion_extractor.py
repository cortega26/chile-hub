"""Extrae permisos de edificación (viviendas) por comuna desde el Centro de
Estudios de Ciudad y Territorio (CEDOC) del MINVU.

Fuente: MINVU CEDOC — repositorio de estadísticas, categoría Permisos de Edificación
URL: https://centrodeestudios.minvu.gob.cl/repositorio/categoria/permisos-de-edificacion
Formato: 1 XLSX ("Viviendas unidades y superficie según año y comuna",
  serie 2002 en adelante) con 6 hojas: número/m2 × {total, casas,
  departamentos}. Grano comuna-año, con comuna y región por fila.

Descubrimiento del archivo: la página del repositorio CEDOC enlaza cada
publicación con su `biblionumber` estable (25583 para este archivo); el
extractor resuelve el href actual y decodifica el parámetro `uri`. Si la
página no responde, usa la URL directa conocida; si la descarga falla,
reutiliza el último snapshot crudo.

Alcance MVP: solo el archivo anual por comuna (el mensual por comuna queda
como extensión documentada). `estado_dato` es "provisional" para años
marcados `(*)` en el encabezado y para el último año de la serie (aún en
curso), "definitivo" en el resto.
"""

import datetime
import os
import re
import sys
import urllib.parse
from pathlib import Path
from typing import Any

import openpyxl
import polars as pl
import requests

UTC = datetime.timezone.utc

ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "../.."))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

try:
    from src.extractors.base import (
        BaseExtractor,
        ensure_staging_directories,
        write_staging_metadata,
    )
except ModuleNotFoundError:
    from base import BaseExtractor, ensure_staging_directories, write_staging_metadata

try:
    from src.extractors.http_utils import stealth_get as _stealth_get
except ModuleNotFoundError:
    from http_utils import stealth_get as _stealth_get

try:
    from src.extractors.region_utils import norm_text
except ModuleNotFoundError:
    from region_utils import norm_text

DATA_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "../../data"))
RAW_DIR = os.path.join(DATA_DIR, "raw")
STAGING_DIR = os.path.join(DATA_DIR, "staging")
STAGING_CSV_PATH = os.path.join(STAGING_DIR, "permisos_edificacion.csv")
METADATA_PATH = os.path.join(STAGING_DIR, "permisos_edificacion.metadata.json")
COMUNAS_CSV_PATH = os.path.join(STAGING_DIR, "comunas.csv")

# ── Fuente ───────────────────────────────────────────────────────────────────
REPOSITORIO_URL = (
    "https://centrodeestudios.minvu.gob.cl/repositorio/categoria/permisos-de-edificacion"
)
# Clave estable de la publicación en el catálogo Koha del MINVU.
BIBLIONUMBER_ANUAL = "25583"
# URL directa conocida (el `id` puede rotar al re-subir el archivo; el
# descubrimiento vía `biblionumber` es la vía primaria).
HARDCODED_XLSX_URL = (
    "https://catalogo.minvu.cl/cgi-bin/koha/opac-retrieve-file.pl"
    "?id=adc446143de3847f7f66d3f93b83ddf3"
)

SOURCE_NAME = "MINVU — Centro de Estudios de Ciudad y Territorio (CEDOC)"
LICENSE_URL = REPOSITORIO_URL

REUSE_POLICY = {
    "status": "open-attribution",
    "license": "Uso autorizado con cita (MINVU CEDOC)",
    "license_url": LICENSE_URL,
    "attribution_required": True,
    "redistribution_ok": True,
    "summary": (
        "Estadísticas de permisos de edificación del CEDOC (MINVU). "
        "El Centro autoriza el uso de la información citando la fuente."
    ),
}

# Hojas del XLSX: (fragmento a matchear, columna canónica).
SHEET_COLUMNS = (
    ("departamentos", "numero", "unidades_departamentos"),
    ("departamentos", "m2", "superficie_m2_departamentos"),
    ("casas", "numero", "unidades_casas"),
    ("casas", "m2", "superficie_m2_casas"),
    ("total", "numero", "unidades_total"),
    ("total", "m2", "superficie_m2_total"),
)

METRIC_COLUMNS = [
    "unidades_total",
    "superficie_m2_total",
    "unidades_casas",
    "superficie_m2_casas",
    "unidades_departamentos",
    "superficie_m2_departamentos",
]

# Grafías históricas de comunas → clean actual (norm_text ya aplicado).
# El archivo escribe "Paiguano" y "Marchihue"; la DPA vigente usa
# "Paihuano" y "Marchigüe".
COMUNA_ALIAS: dict[str, str] = {
    "paiguano": "paihuano",
    "marchihue": "marchigue",
}

# ── Fallback mínimo ─────────────────────────────────────────────────────────
# Estructural, no cobertura: evita que el build se rompa si el CEDOC no
# responde. El job `publish` de CI rechaza datasets en fallback (§5).
FALLBACK_ROWS = [
    {
        "anio": 2023,
        "codigo_region": "13",
        "codigo_comuna": "13101",
        "nombre_comuna": "Santiago",
        "unidades_total": 0,
        "superficie_m2_total": 0,
        "unidades_casas": 0,
        "superficie_m2_casas": 0,
        "unidades_departamentos": 0,
        "superficie_m2_departamentos": 0,
        "estado_dato": "definitivo",
        "fuente": f"{SOURCE_NAME} (fallback)",
        "url_fuente": REPOSITORIO_URL,
        "fecha_fuente": "",
    },
    {
        "anio": 2023,
        "codigo_region": "08",
        "codigo_comuna": "08101",
        "nombre_comuna": "Concepción",
        "unidades_total": 0,
        "superficie_m2_total": 0,
        "unidades_casas": 0,
        "superficie_m2_casas": 0,
        "unidades_departamentos": 0,
        "superficie_m2_departamentos": 0,
        "estado_dato": "definitivo",
        "fuente": f"{SOURCE_NAME} (fallback)",
        "url_fuente": REPOSITORIO_URL,
        "fecha_fuente": "",
    },
]

REQUIRED_COLUMNS = [
    "anio",
    "codigo_region",
    "codigo_comuna",
    "nombre_comuna",
    *METRIC_COLUMNS,
    "estado_dato",
    "fuente",
    "url_fuente",
    "fecha_fuente",
]


# ── Helpers ──────────────────────────────────────────────────────────────────


def _snapshot_path() -> Path:
    """Ruta de snapshot crudo con timestamp."""
    stamp = datetime.datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
    return Path(RAW_DIR) / f"minvu_permisos_edificacion_anual_{stamp}.xlsx"


def _discover_xlsx_url() -> tuple[str, str]:
    """Resuelve la URL actual del XLSX anual. Retorna (url, método).

    Vía primaria: href con `biblionumber=25583` en la página del repositorio
    CEDOC (se decodifica el parámetro `uri`). Secundaria: URL directa
    conocida (el `id` de Koha puede rotar entre versiones).
    """
    try:
        response = _stealth_get(REPOSITORIO_URL, timeout=60)
        response.raise_for_status()
        html = response.text
        for href in re.findall(
            r'href="([^"]*biblionumber=' + BIBLIONUMBER_ANUAL + r'[^"]*)"', html
        ):
            parsed = urllib.parse.urlparse(href.replace("&amp;", "&"))
            params = urllib.parse.parse_qs(parsed.query)
            if "uri" in params and params["uri"]:
                return params["uri"][0], "repositorio-biblionumber"
            if href.startswith("http"):
                return href, "repositorio-tracklinks"
    except (requests.RequestException, ValueError) as exc:
        return HARDCODED_XLSX_URL, f"directa-conocida (descubrimiento falló: {exc})"
    return HARDCODED_XLSX_URL, "directa-conocida (sin match biblionumber en página)"


def _download_xlsx(url: str) -> Path:
    """Descarga el XLSX y guarda snapshot crudo. Retorna la ruta."""
    target = _snapshot_path()
    response = _stealth_get(url, timeout=120)
    response.raise_for_status()
    target.write_bytes(response.content)
    return target


def _load_comunas_lookup() -> dict[str, tuple[str, str, str]]:
    """``nombre_comuna_clean -> (codigo_comuna, codigo_region, nombre_comuna)``."""
    if not os.path.exists(COMUNAS_CSV_PATH):
        raise FileNotFoundError(
            "No se encuentra data/staging/comunas.csv. "
            "Corre primero: python src/extractors/subdere_extractor.py"
        )
    df = pl.read_csv(
        COMUNAS_CSV_PATH,
        schema_overrides={"codigo_comuna": pl.String, "codigo_region": pl.String},
    )
    return {
        row["nombre_comuna_clean"]: (
            row["codigo_comuna"],
            row["codigo_region"],
            row["nombre_comuna"],
        )
        for row in df.iter_rows(named=True)
    }


def _clean_label(value) -> str:
    """Normaliza un nombre de comuna del XLSX para el cruce a CUT."""
    text = str(value).strip()
    # "Cabo de Hornos (Ex - Navarino)" -> "Cabo de Hornos"; "Ñuble (**)" no
    # llega aquí (filas de región), pero la limpieza es la misma.
    text = re.sub(r"\s*\(.*?\)\s*", "", text).strip()
    # Apóstrofes tipográficos ("O’Higgins") -> ASCII.
    text = text.replace("’", "'").replace("‘", "'").replace("´", "'").replace("`", "'")
    clean = norm_text(text)
    return COMUNA_ALIAS.get(clean, clean)


def _to_int(value) -> int:
    """Celda numérica a int; None/'-'/' ' → 0."""
    if value is None:
        return 0
    if isinstance(value, str):
        value = value.strip().replace(".", "").replace(",", "")
        if value in ("", "-"):
            return 0
    try:
        return int(float(value))
    except (ValueError, TypeError):
        return 0


def _year_from_header(cell) -> tuple[int | None, bool]:
    """Extrae (anio, es_provisional) de un encabezado de columna.

    "2024 (*)" → (2024, True); "2018 (***)" → (2018, False): el triple
    asterisco es nota metodológica, no provisionalidad. 2002 (int) →
    (2002, False).
    """
    if cell is None:
        return None, False
    text = str(cell).strip()
    digits = re.sub(r"\D", "", text)
    if not digits:
        return None, False
    provisional = bool(re.search(r"\(\*\)(?!\*)", text))
    return int(digits), provisional


def _sheet_target(sheet_name: str) -> str | None:
    """Mapea el nombre de hoja a columna canónica (o None si no aplica)."""
    n = norm_text(sheet_name)
    if "departamentos" in n:
        tipo = "departamentos"
    elif "casas" in n:
        tipo = "casas"
    elif "total" in n:
        tipo = "total"
    else:
        return None
    metrica = "m2" if "m2" in n else "numero"
    for want_tipo, want_metrica, column in SHEET_COLUMNS:
        if tipo == want_tipo and metrica == want_metrica:
            return column
    return None


def _parse_sheet(
    rows: list[list], target: str
) -> tuple[dict[str, dict[int, int]], dict[int, int], list[int]]:
    """Parsea una hoja a {clean_comuna: {anio: valor}}.

    Retorna (valores, totales_pais, anios_provisionales). Las filas de región
    y "Total País" no entran a valores: el total se retorna para
    reconciliación. Filas duplicadas de comuna (Ñuble bajo "Biobío" y bajo
    "Ñuble (ex-Biobío)") se SUMAN: la serie se parte en 2018 y cada sección
    trae ceros donde la otra trae datos.
    """
    # Fila de encabezado: primera con "Región" en col A y años en col C+.
    header_idx = next(
        i
        for i, r in enumerate(rows[:12])
        if r
        and isinstance(r[0], str)
        and norm_text(r[0]) == "region"
        and r[1] is not None
        and norm_text(str(r[1])) == "comuna"
    )
    header = rows[header_idx]
    year_cols: list[tuple[int, int, bool]] = []  # (índice col, anio, provisional*)
    for j in range(2, len(header)):
        year, star = _year_from_header(header[j])
        if year is not None:
            year_cols.append((j, year, star))
    if not year_cols:
        raise KeyError(f"hoja sin columnas de año: {target}")
    max_year = max(y for _, y, _ in year_cols)

    valores: dict[str, dict[int, int]] = {}
    totales: dict[int, int] = {}
    for r in rows[header_idx + 1 :]:
        if not r:
            continue
        a = r[0] if len(r) > 0 else None
        b = r[1] if len(r) > 1 else None
        label_a = str(a).strip() if a is not None and str(a).strip() else ""
        if not label_a:
            continue
        if norm_text(label_a) in ("total pais", "total país"):
            for j, year, _ in year_cols:
                totales[year] = totales.get(year, 0) + _to_int(r[j] if j < len(r) else None)
            continue
        if b is None or not str(b).strip():
            continue  # fila de región u otra agregación: se omite
        clean = _clean_label(b)
        bucket = valores.setdefault(clean, {})
        for j, year, _ in year_cols:
            bucket[year] = bucket.get(year, 0) + _to_int(r[j] if j < len(r) else None)

    provisionales = sorted({y for _, y, star in year_cols if star or y >= max_year})
    # El último año de la serie siempre está en curso al publicarse.
    return valores, totales, provisionales


def _gate_reconciliacion_anual(
    rows: list[dict[str, Any]], totals: dict[tuple[str, int], int]
) -> tuple[list[dict[str, Any]], list[str]]:
    """Descarta años completos que no reconcilian con la fila Total País.

    El validador exige cobertura total por año y aborta ruidosamente, en vez
    de publicar métricas corruptas por un cambio de layout.
    Retorna (filas_conservadas, mensajes_error).
    """
    anios = sorted({r["anio"] for r in rows})
    malos: set[int] = set()
    errores: list[str] = []
    for target in METRIC_COLUMNS:
        for year in anios:
            esperado = totals.get((target, year))
            obtenido = sum(r[target] for r in rows if r["anio"] == year)
            if esperado is not None and esperado != obtenido:
                if year not in malos:
                    malos.add(year)
                errores.append(
                    f"ERROR reconciliación {target} {year}: suma comunas={obtenido} "
                    f"vs Total País={esperado}; año omitido"
                )
    kept = [r for r in rows if r["anio"] not in malos]
    return kept, errores


def fetch_data() -> tuple[list[dict], str, str, list[str]]:
    """Obtiene permisos de edificación por comuna desde MINVU CEDOC.

    Retorna (rows, source_mode, source_url, notes).
    """
    ensure_staging_directories()
    notes: list[str] = []
    fecha_fuente = datetime.datetime.now(UTC).strftime("%Y-%m-%d")

    try:
        lookup = _load_comunas_lookup()
    except FileNotFoundError as exc:
        return FALLBACK_ROWS, "fallback", REPOSITORIO_URL, [str(exc)]

    try:
        url, metodo = _discover_xlsx_url()
        notes.append(f"descubrimiento: {metodo}")
    except (requests.RequestException, ValueError) as exc:
        notes.append(f"descubrimiento falló ({exc}); usando URL directa conocida")
        url = HARDCODED_XLSX_URL

    try:
        try:
            path = _download_xlsx(url)
        except (requests.RequestException, OSError) as exc:
            snapshots = sorted(Path(RAW_DIR).glob("minvu_permisos_edificacion_anual_*.xlsx"))
            if not snapshots:
                raise
            path = snapshots[-1]
            notes.append(f"descarga falló, usando snapshot {path.name} ({exc})")
    except (requests.RequestException, OSError) as exc:
        notes.append(f"sin datos live ni snapshot ({exc})")
        return FALLBACK_ROWS, "fallback", REPOSITORIO_URL, notes

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_names = wb.sheetnames
    finally:
        wb.close()

    # Acumular métricas por (comuna, anio) a través de las 6 hojas.
    series: dict[tuple[str, int], dict[str, int]] = {}
    totals: dict[tuple[str, int], int] = {}
    provisionales: set[int] = set()
    hojas_ok = []
    for name in sheet_names:
        target = _sheet_target(name)
        if target is None:
            notes.append(f"hoja '{name}' no mapeada a métrica, se omite")
            continue
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            sheet_rows = list(wb[name].iter_rows(values_only=True))
        finally:
            wb.close()
        try:
            valores, tot, prov = _parse_sheet(sheet_rows, target)
        except (KeyError, ValueError, IndexError, StopIteration) as exc:
            notes.append(f"hoja '{name}': omitida por error de parseo ({exc})")
            continue
        hojas_ok.append(name)
        provisionales.update(prov)
        for clean, by_year in valores.items():
            for year, value in by_year.items():
                series.setdefault((clean, year), {})[target] = value
        for year, value in tot.items():
            totals[(target, year)] = value

    if not series:
        notes.append("ninguna hoja aportó filas")
        return FALLBACK_ROWS, "fallback", REPOSITORIO_URL, notes

    unmatched = sorted(c for (c, _) in series if c not in lookup)
    if unmatched:
        notes.append(f"{len(unmatched)} comunas sin match CUT (se omiten): {unmatched[:10]}")

    rows: list[dict[str, Any]] = []
    for (clean, year), metrics in sorted(series.items()):
        hit = lookup.get(clean)
        if hit is None:
            continue
        codigo_comuna, codigo_region, nombre_comuna = hit
        row = {
            "anio": year,
            "codigo_region": codigo_region,
            "codigo_comuna": codigo_comuna,
            "nombre_comuna": nombre_comuna,
            "estado_dato": "provisional" if year in provisionales else "definitivo",
            "fuente": SOURCE_NAME,
            "url_fuente": url,
            "fecha_fuente": fecha_fuente,
        }
        for col in METRIC_COLUMNS:
            row[col] = metrics.get(col, 0)
        rows.append(row)

    # Reconciliación contra filas "Total País" del archivo. Un año que no
    # cuadra se descarta completo (el validador exige cobertura total por
    # año y aborta ruidosamente) en vez de publicar métricas corruptas por
    # un cambio de layout.
    rows, gate_errors = _gate_reconciliacion_anual(rows, totals)
    notes.extend(gate_errors)
    notes.append(
        f"{len(rows)} filas desde '{path.name}' "
        f"(hojas: {', '.join(hojas_ok)}; años provisionales: {sorted(provisionales) or 'ninguno'})"
    )
    return rows, "live", REPOSITORIO_URL, notes


def normalize_rows(rows: list[dict]) -> pl.DataFrame:
    """Convierte filas a DataFrame canónico."""
    if not rows:
        return pl.DataFrame(
            schema={col: pl.String for col in REQUIRED_COLUMNS}
            | {"anio": pl.Int64}
            | {col: pl.Int64 for col in METRIC_COLUMNS}
        )

    df = pl.DataFrame(rows, strict=False)
    for col in REQUIRED_COLUMNS:
        if col not in df.columns:
            df = df.with_columns(pl.lit(None).alias(col))

    return df.with_columns(
        pl.col("anio").cast(pl.Int64),
        pl.col("codigo_region").cast(pl.String),
        pl.col("codigo_comuna").cast(pl.String),
        pl.col("nombre_comuna").cast(pl.String),
        *[pl.col(col).cast(pl.Int64) for col in METRIC_COLUMNS],
        pl.col("estado_dato").cast(pl.String),
        pl.col("fuente").cast(pl.String),
        pl.col("url_fuente").cast(pl.String),
        pl.col("fecha_fuente").cast(pl.String),
    ).select(REQUIRED_COLUMNS)


def build_metadata(mode: str, source_url: str, notes: list[str], row_count: int) -> dict:
    """Construye metadatos estándar para el dataset."""
    return {
        "dataset": "permisos_edificacion",
        "source_name": SOURCE_NAME,
        "source_url": source_url,
        "source_mode": mode,
        "source_detail": (
            "Viviendas en unidades y superficie (m2) por comuna y año, serie "
            "desde 2002 (MINVU CEDOC, en base a permisos otorgados por las "
            "Direcciones de Obras Municipales e INE)."
        ),
        "refreshed_at_utc": datetime.datetime.now(UTC).isoformat(),
        "record_count": row_count,
        "fields": REQUIRED_COLUMNS,
        "notes": notes,
        "reuse_policy": REUSE_POLICY,
    }


def process_permisos_edificacion() -> dict:
    """Ejecuta el flujo completo de extracción y staging."""
    rows, mode, source_url, notes = fetch_data()
    df = normalize_rows(rows)

    metadata = build_metadata(mode, source_url, notes, df.height)

    ensure_staging_directories()
    df.write_csv(STAGING_CSV_PATH)
    write_staging_metadata(METADATA_PATH, metadata)
    print(f"permisos_edificacion: {df.height} filas escritas en staging (mode={mode})")
    if mode == "fallback":
        # Diagnóstico visible en el log de CI: sin esto, un fallback por 403/WAF
        # sólo se ve como "2 filas (mode=fallback)" y la causa queda enterrada
        # en el metadata de staging (que no se publica).
        print(f"permisos_edificacion notas: {' | '.join(notes)}")

    return metadata


# ── Extractor ────────────────────────────────────────────────────────────────


class PermisosEdificacionExtractor(BaseExtractor):
    """Extractor de permisos de edificación por comuna (MINVU CEDOC)."""

    @property
    def dataset_name(self) -> str:
        return "permisos_edificacion"

    def fetch(self, **kwargs):
        return fetch_data()

    def normalize(self, raw_data):
        rows, _mode, _url, _notes = raw_data
        return normalize_rows(rows)

    def validate(self, df, metadata: dict) -> dict:
        from src.validation import validate_permisos_edificacion

        return validate_permisos_edificacion(df, metadata)

    def write_staging(self, df, metadata: dict) -> Path:
        ensure_staging_directories()
        output = Path(STAGING_CSV_PATH)
        df.write_csv(str(output))
        merged = {
            **metadata,
            "dataset": self.dataset_name,
            "refreshed_at_utc": datetime.datetime.now(UTC).isoformat(),
            "record_count": df.height,
        }
        write_staging_metadata(METADATA_PATH, merged)
        return output


if __name__ == "__main__":
    process_permisos_edificacion()
